import pandas as pd
import docx
from fnreadfilename import readfilesname
path = '10Mar'
x = readfilesname(path)
print(x)
a=0
while a<(len(x)):
    df1 = pd.read_csv(path+'\\'+x[a],header=None)
    sorted_df = df1.sort_values(by=[0], ascending=False)
    df2 = sorted_df[[1,2,3,4,5,7,9]]
    header = ['DB Name','State','Recovery model','Total size(MB)','Data size(MB)','Log size(MB)','Full last date']
    df = df2.rename(columns={1: 'DB Name',2:'State',3:'Recovery model',4:'Total size(MB)',5:'Data size(MB)',7:'Log size(MB)',9:'Full last date'})
    doc = docx.Document()
    # Initialise the table
    t = doc.add_table(rows=1, cols=df.shape[1])
    # Add borders
    t.style = 'TableGrid'
    # Add the column headings
    for j in range(df.shape[1]):
        t.cell(0, j).text = df.columns[j]
    # Add the body of the data frame
    for i in range(df.shape[0]):
        row = t.add_row()
        for j in range(df.shape[1]):
            cell = df.iat[i, j]
            row.cells[j].text = str(cell)
    # Save the Word doc
    doc.save('wordfile\\'+path+x[a]+'.docx')
    a=a+1



